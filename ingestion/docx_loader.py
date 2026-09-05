"""
DOCX document loader using python-docx.
Extracts text from paragraphs and tables in Microsoft Word files.
"""

import docx
from typing import Dict, Any, Union
import io
import os


def load_docx(file_input: Union[str, bytes, io.BytesIO], filename: str = "document.docx") -> Dict[str, Any]:
    """
    Extract text content from a DOCX file using python-docx.

    Args:
        file_input: File path (str), raw bytes, or BytesIO.
        filename: Name of the DOCX document.

    Returns:
        Dict[str, Any]: {
            "text": Extracted document text,
            "source": filename,
            "file_type": "docx"
        }
    """
    try:
        if isinstance(file_input, str):
            doc = docx.Document(file_input)
        elif isinstance(file_input, bytes):
            doc = docx.Document(io.BytesIO(file_input))
        elif hasattr(file_input, "read"):
            content = file_input.read()
            if hasattr(file_input, "seek"):
                file_input.seek(0)
            doc = docx.Document(io.BytesIO(content))
        else:
            raise ValueError("Unsupported input type for DOCX loader.")

        elements = []

        # Extract text from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                elements.append(para.text.strip())

        # Extract text from tables
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_data:
                    table_rows.append(" | ".join(row_data))
            if table_rows:
                elements.append("--- Table Data ---\n" + "\n".join(table_rows))

        full_text = "\n\n".join(elements)

        if not full_text.strip():
            raise ValueError(f"DOCX file '{filename}' contains no extractable text.")

        return {
            "text": full_text,
            "source": os.path.basename(filename),
            "file_type": "docx"
        }
    except Exception as e:
        raise RuntimeError(f"Failed to extract text from DOCX '{filename}': {str(e)}")
